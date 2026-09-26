"""Content scanning: what is found, what each action does, and what is never written down.

Every credential below is generated at run time, by concatenation, from
random characters: none is a real secret, and none is committed in a form a
secret scanner would (rightly) refuse to push.
"""

from __future__ import annotations

import asyncio
import base64
import copy
import hashlib
import json
import random
import string
from pathlib import Path

import pytest
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
from fastapi.testclient import TestClient

from backend.agents.orchestrator import AgentOrchestrator, _dlp_holds, _model_text
from backend.core.config import PROJECT_ROOT, ConfigError, get_config
from backend.core.schemas import EvidenceItem, PolicyDecision, Sensitivity, TaskStatus
from backend.policy.gateway import get_policy_gateway
from backend.security import dlp
from tests.test_usage_telemetry import STREAM_STATS, USER, FakeClient, blocking_result, make_task, orchestrator_with

ROOT = Path(__file__).resolve().parents[2]
_rng = random.Random(20260925)


def _chars(alphabet: str, length: int) -> str:
    return "".join(_rng.choice(alphabet) for _ in range(length))


def aws_key() -> str:
    return "AK" + "IA" + _chars(string.ascii_uppercase + "234567", 16)


def github_token() -> str:
    return "gh" + "p_" + _chars(string.ascii_letters + string.digits, 36)


def slack_token() -> str:
    return "xo" + "xb-" + _chars(string.digits, 11) + "-" + _chars(string.ascii_letters + string.digits, 24)


def jwt() -> str:
    def part(payload: dict) -> str:
        return base64.urlsafe_b64encode(json.dumps(payload).encode()).decode().rstrip("=")

    return f"{part({'alg': 'HS256', 'typ': 'JWT'})}.{part({'sub': 'svc-historian', 'iat': 1758800000})}." + _chars(
        string.ascii_letters + string.digits + "-_", 43
    )


def private_key() -> str:
    return Ed25519PrivateKey.generate().private_bytes(
        serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()
    ).decode()


def aadhaar(valid: bool = True) -> str:
    """A 12-digit number whose Verhoeff check digit is right, or deliberately wrong."""
    stem = str(_rng.randint(2, 9)) + _chars(string.digits, 10)
    good = next(d for d in string.digits if dlp.verhoeff_valid(stem + d))
    last = good if valid else str((int(good) + 1) % 10)
    number = stem + last
    return f"{number[:4]} {number[4:8]} {number[8:]}"


def kinds(text: str, boundary: str = "prompt") -> list[str]:
    return [finding.detector for finding in dlp.scan(text, boundary).findings]


# ------------------------------------------------------------------ detectors
@pytest.mark.parametrize(
    ("make", "detector"),
    [
        (lambda: f"Historian service key:\n{private_key()}", "private_key"),
        (lambda: f"aws access key id {aws_key()} was left in the export.", "aws_access_key"),
        (lambda: f"clone with {github_token()} before the shutdown", "github_token"),
        (lambda: f"alerts posted with {slack_token()}", "slack_token"),
        (lambda: f"session cookie {jwt()}", "jwt"),
        (lambda: "Authorization: Bearer " + _chars(string.ascii_letters + string.digits, 40), "bearer_token"),
        (lambda: "historian_db_password = Zq8" + _chars(string.ascii_letters, 9), "credential_assignment"),
        (lambda: 'opc_ua:\n  api_key: "' + _chars(string.ascii_letters + string.digits, 24) + '"', "credential_assignment"),
        (lambda: "Password: Tr0" + _chars(string.ascii_lowercase, 6) + "!", "credential_assignment"),
        (lambda: "DSN postgres://historian:" + _chars(string.ascii_letters, 6) + "7@10.4.2.19:5432/plant", "connection_string_password"),
        (lambda: f"Operator Aadhaar {aadhaar()} recorded on the permit.", "aadhaar"),
        (lambda: "Contractor PAN ABCPK4821M on the gate pass.", "pan"),
        (lambda: "Raised by r.sharma@refinery-ops.co.in on shift B.", "email"),
        (lambda: "Call the shift lead on +91 98450 12345 if the alarm persists.", "phone_in"),
        (lambda: "**Classification:** Restricted — unreleased design information", "classification_marker"),
        (lambda: "SECRET\nTurbine trip logic, rev 3.", "classification_marker"),
    ],
)
def test_each_detector_finds_what_it_is_for(make, detector: str) -> None:
    assert detector in kinds(make())


@pytest.mark.parametrize(
    "text",
    [
        # Checksums: a 12-digit number that fails Verhoeff is not an Aadhaar number.
        f"Work order {aadhaar(valid=False)} closed.",
        # Structure: the fourth letter of a PAN is the holder type; D is none.
        "Reference ABCDE1234F",
        # The values documentation uses as examples.
        "aws_access_key_id AKIAIOSFODNN7EXAMPLE",
        "-----BEGIN RSA PRIVATE KEY-----\n...\n-----END RSA PRIVATE KEY-----",
        "postgres://user:password@localhost/db",
        "Contact admin@example.com or ops@plant.test",
        "Helpline 9876543210",
        "api_key = <your-api-key>",
        "password = ${HISTORIAN_PASSWORD}",
        "token = xxxxxxxxxxxxxxxx",
        # Procedure prose and code are not credentials.
        "Password: required at every login.",
        "def get_current_user(token: SessionToken) -> User:",
        "max_tokens=1200000",
        '{"username": username, "password": self.password}',
        # A JWT-shaped string whose header is not JSON.
        "eyJub3RqdXN0.eyJzdWJqZWN0MTIz.c2lnbmF0dXJlMTIz",
        # "restricted" as a word, or a banner that is not in capitals.
        "Access to the tank farm is restricted to permit holders.",
        "restricted",
        # Engineering figures are not phone numbers.
        "Shell course 2 measured 9.8450 mm at grid 12345; tag 7000012345.5",
    ],
)
def test_placeholders_invalid_checksums_and_prose_are_not_findings(text: str) -> None:
    assert kinds(text) == []


def test_the_demonstration_corpus_carries_only_its_own_markings() -> None:
    texts = [p.read_text(errors="ignore") for p in (ROOT / "sample_data").rglob("*")
             if p.suffix in (".md", ".csv", ".txt", ".json")]
    found = {d for text in texts for d in kinds(text)}
    assert texts and found <= {"classification_marker"}


def test_a_marking_raises_to_the_level_it_asserts() -> None:
    assert dlp.scan("**Classification:** Sensitive — incident investigation", "upload").floor == "sensitive"
    assert dlp.scan("TOP SECRET\nnothing else", "upload").floor == "restricted"
    assert dlp.scan("Classification: Normal", "upload").floor == "normal"


def test_the_verhoeff_check_is_the_published_one() -> None:
    # The worked example from the algorithm's description: 236 -> check digit 3.
    assert dlp.verhoeff_valid("2363") and not dlp.verhoeff_valid("2364")


# ---------------------------------------------------------------- redaction
def test_redaction_names_the_detector_and_keeps_nothing_of_the_value() -> None:
    key = aws_key()
    text = f"export key {key} before the audit"
    result = dlp.scan(text, "prompt")
    shown = dlp.redact(text, result)
    assert key not in shown and key[4:] not in shown
    assert f"[redacted: aws_access_key #{result.findings[0].fingerprint[:8]}]" in shown
    assert shown.startswith("export key ") and shown.endswith(" before the audit")


def test_overlapping_findings_are_redacted_as_one() -> None:
    secret = "Zq8" + _chars(string.ascii_letters, 9)
    text = f"mongodb://svc:{secret}@historian.refinery-ops.co.in/db"
    shown = dlp.redact(text, dlp.scan(text, "prompt"), actions=("redact", "allow"))
    assert secret not in shown and "@" not in shown.split("svc:")[1][:12]


def test_fingerprints_are_keyed_stable_and_not_a_plain_hash() -> None:
    number = aadhaar()
    first = dlp.scan(f"A {number}", "prompt").findings[0].fingerprint
    again = dlp.scan(f"Another record, {number}.", "prompt").findings[0].fingerprint
    assert first == again
    digits = number.replace(" ", "")
    assert first not in {hashlib.sha256(v.encode()).hexdigest()[:16] for v in (digits, number, f"aadhaar:{digits}")}
    assert (get_config().settings.storage_root / "keys" / "dlp-fingerprint.key").stat().st_size == 32


def test_a_structure_is_redacted_string_by_string() -> None:
    number = aadhaar()
    content, found = dlp.redact_value(
        {"title": "Permit", "sections": [{"heading": "Holder", "body": f"Aadhaar {number}"}], "rev": 3}, "deliverable"
    )
    assert number not in json.dumps(content) and content["rev"] == 3
    assert [f.detector for f in found] == ["aadhaar"]


def test_the_stream_never_shows_a_value_the_answer_would_redact() -> None:
    key = aws_key()
    partial = f"The export used {key[:10]}"
    assert key[:10] not in dlp.mask_stream(partial, final=False)
    number = aadhaar()
    assert number[:9] not in dlp.mask_stream(f"Operator {number[:10]}", final=False)
    finished = dlp.mask_stream(f"The export used {key} yesterday.", final=True)
    assert key not in finished and "[redacted: aws_access_key" in finished


# ------------------------------------------------------------------- policy
@pytest.fixture
def policy_override(monkeypatch):
    """Rewrite policies/dlp.yaml for one test, through the same validation."""
    config = get_config()
    original = copy.deepcopy(config.dlp)

    def apply(mutate) -> None:
        changed = copy.deepcopy(original)
        mutate(changed)
        monkeypatch.setattr(config, "dlp", changed)
        dlp._policy_cache = None
        dlp.policy()

    yield apply
    monkeypatch.setattr(config, "dlp", original)
    dlp._policy_cache = None


def test_the_policy_refuses_an_action_a_boundary_cannot_take(policy_override) -> None:
    def redact_uploads(p):
        p["detectors"]["aws_access_key"]["actions"]["upload"] = "redact"

    with pytest.raises(ConfigError, match="upload"):
        policy_override(redact_uploads)


def test_the_policy_must_name_every_detector_and_no_other(policy_override) -> None:
    with pytest.raises(ConfigError, match="missing"):
        policy_override(lambda p: p["detectors"].pop("pan"))
    with pytest.raises(ConfigError, match="unknown"):
        policy_override(lambda p: p["detectors"].__setitem__("iban", dict(p["detectors"]["pan"])))


def test_actions_come_from_the_policy_file(policy_override) -> None:
    key = aws_key()
    assert dlp.scan(key, "answer").action == "redact"
    policy_override(lambda p: p["detectors"]["aws_access_key"]["actions"].__setitem__("answer", "require_approval"))
    assert dlp.scan(key, "answer").action == "require_approval"


# ------------------------------------------------------------ the boundaries
def _item(excerpt: str, **extra) -> EvidenceItem:
    fields = dict(id="F1", source_document="historian-export.txt", excerpt=excerpt, kind="uploaded_file",
                  classification=Sensitivity.CONFIDENTIAL)
    fields.update(extra)
    return EvidenceItem(**fields)


def _audit_text() -> str:
    path = Path(str(get_config().settings.audit.get("log_file")))
    path = path if path.is_absolute() else PROJECT_ROOT / path
    return path.read_text()


def test_prompt_redact_the_model_sees_a_marker_and_the_item_is_raised_and_recorded() -> None:
    key = aws_key()
    task, orchestrator = make_task(evidence=[_item(f"Export header. Key {key}. Rows follow.")]), AgentOrchestrator()
    assert asyncio.run(orchestrator._scan_evidence(task, USER)) == 1
    item = task.evidence[0]
    assert item.excerpt.endswith("Rows follow.") and key in item.excerpt  # kept whole for the reviewer
    assert key not in _model_text(item) and "[redacted: aws_access_key" in _model_text(item)
    assert item.classification == Sensitivity.RESTRICTED
    event = task.policy_events[-1]
    assert event.action == "dlp.prompt" and event.decision == PolicyDecision.ALLOW and key not in event.reason
    assert event.rule == "dlp.yaml:detectors.aws_access_key.actions.prompt"
    # Scanning again records nothing new.
    assert asyncio.run(orchestrator._scan_evidence(task, USER)) == 0
    assert key not in _audit_text()


def test_a_marking_is_recorded_only_when_it_says_more_than_the_label() -> None:
    header = "**Classification:** Confidential — Internal Use Only\n\nInspection intervals follow."
    task = make_task(evidence=[_item(header, id="S1", kind="knowledge_base")])
    assert asyncio.run(AgentOrchestrator()._scan_evidence(task, USER)) == 0 and task.policy_events == []

    memo = "**Classification:** Restricted — unreleased design information"
    task = make_task(evidence=[_item(memo, id="S1", kind="knowledge_base")])
    assert asyncio.run(AgentOrchestrator()._scan_evidence(task, USER)) == 1
    assert task.evidence[0].classification == Sensitivity.RESTRICTED
    assert "(was confidential)" in task.policy_events[-1].reason


def test_prompt_block_withholds_the_whole_item() -> None:
    key = private_key()
    item = _item(f"Service account:\n{key}")
    shown = _model_text(item)
    assert "BEGIN" not in shown and "withheld from the model by content scanning" in shown and "private_key" in shown


def test_prompt_allow_shows_the_value_and_records_it() -> None:
    task = make_task(evidence=[_item("Raised by r.sharma@refinery-ops.co.in on shift B.")])
    asyncio.run(AgentOrchestrator()._scan_evidence(task, USER))
    assert "r.sharma@refinery-ops.co.in" in _model_text(task.evidence[0])
    assert task.policy_events[-1].decision == PolicyDecision.ALLOW
    assert task.evidence[0].extraction_data["dlp"][0]["detector"] == "email"


def test_prompt_require_approval_holds_the_run(policy_override) -> None:
    policy_override(lambda p: p["detectors"]["email"]["actions"].__setitem__("prompt", "require_approval"))
    task = make_task(evidence=[_item("Raised by r.sharma@refinery-ops.co.in on shift B.")])
    asyncio.run(AgentOrchestrator()._scan_evidence(task, USER))
    assert task.policy_events[-1].decision == PolicyDecision.REQUIRE_APPROVAL and _dlp_holds(task) == 1
    required, reasons, _ = get_policy_gateway().approval_requirement(
        task.profile, prompt=task.prompt, dlp_findings=_dlp_holds(task))
    assert required and any(reason.startswith("sensitive_content") for reason in reasons)


def test_answer_redact_and_block(policy_override) -> None:
    key = aws_key()
    task, orchestrator = make_task(), AgentOrchestrator()
    released = asyncio.run(orchestrator._release_text(task, USER, f"The export used {key}."))
    assert key not in released and "[redacted: aws_access_key" in released and _dlp_holds(task) == 0

    policy_override(lambda p: p["detectors"]["aws_access_key"]["actions"].__setitem__("answer", "block"))
    task = make_task()
    released = asyncio.run(orchestrator._release_text(task, USER, f"The export used {key}."))
    assert key not in released and released.startswith("[The answer was withheld by content scanning")
    assert task.policy_events[-1].decision == PolicyDecision.DENY and _dlp_holds(task) == 1
    assert key not in _audit_text()


def test_deliverable_require_approval_redact_and_block() -> None:
    orchestrator = AgentOrchestrator()
    limitations: list[str] = []

    task = make_task()
    content = {"title": "Permit", "summary": "Issued to r.sharma@refinery-ops.co.in."}
    released, _ = asyncio.run(orchestrator._scan_deliverable(task, USER, content, [], limitations))
    assert released == content and task.policy_events[-1].decision == PolicyDecision.REQUIRE_APPROVAL
    assert _dlp_holds(task) == 1

    number = aadhaar()
    task = make_task()
    released, _ = asyncio.run(orchestrator._scan_deliverable(
        task, USER, {"title": "Permit", "summary": "Holder"}, [_item(f"Aadhaar {number}")], limitations))
    assert released is not None and task.policy_events[-1].decision == PolicyDecision.ALLOW
    quoted = orchestrator._deliverable_copy({"title": "x"}, [_item(f"Aadhaar {number}")])[1]
    assert number not in quoted[0].excerpt

    task = make_task()
    released, floor = asyncio.run(orchestrator._scan_deliverable(
        task, USER, {"title": "Keys", "summary": private_key()}, [], limitations))
    assert released is None and floor is None
    assert task.policy_events[-1].decision == PolicyDecision.DENY and _dlp_holds(task) == 1
    assert any(line.startswith("No deliverable was rendered") for line in limitations)


def test_the_render_itself_refuses_blocked_evidence() -> None:
    from backend.tools.registry import ToolContext

    orchestrator = AgentOrchestrator()
    task = make_task()
    task.profile = task.profile.model_copy(update={"produces_deliverable": True, "deliverable_format": "md"})
    context = ToolContext(user=USER, task_id=task.id, sensitivity=Sensitivity.NORMAL, files=[],
                          workspace=get_config().settings.storage_root / "workspaces")
    asyncio.run(orchestrator._render_deliverable(
        task, context, {"title": "Keys", "summary": "See evidence."}, [_item(private_key())], []))
    assert task.deliverables == [] and task.policy_events[-1].action == "dlp.deliverable"
    assert task.policy_events[-1].decision == PolicyDecision.DENY


# ----------------------------------------------------------- the whole run
@pytest.mark.asyncio
async def test_a_run_never_shows_the_model_or_the_reader_a_redacted_value() -> None:
    key = aws_key()
    # A model that repeats the key anyway -- as if it had seen it elsewhere.
    client = FakeClient(
        fragments=["The export ", f"used {key[:9]}", f"{key[9:]} ", "[F1]."],
        stats=STREAM_STATS,
        results=[blocking_result(text='{"calculations": []}', prompt_eval_count=10, eval_count=5)],
    )
    orchestrator, recorder, _ = orchestrator_with(client)
    task = make_task(evidence=[_item(f"Export header. Key {key}. Rows follow.", kind="knowledge_base")])

    finished = await orchestrator.run(task, USER)

    assert all(key not in call["prompt"] for call in client.stream_calls)
    assert all(key[:9] not in frame["text"] for frame in recorder.named("task.token"))
    assert key not in finished.answer and key not in recorder.named("task.answer")[-1]["answer"]
    # The evidence carried a restricted value, so the run is restricted and held.
    assert finished.profile.sensitivity == Sensitivity.RESTRICTED
    assert finished.status == TaskStatus.AWAITING_APPROVAL
    assert any(line.startswith("Content scanning:") for line in finished.verification.limitations)
    assert key not in _audit_text()


# ------------------------------------------------------------------ uploads
@pytest.fixture(scope="module")
def client():
    from backend.api.main import create_app

    with TestClient(create_app()) as test_client:
        token = test_client.post("/api/auth/login", json={"username": "engineer", "password": "workbench"}).json()["token"]
        test_client.headers["Authorization"] = f"Bearer {token}"
        yield test_client


def test_upload_block_refuses_a_private_key_and_stores_nothing(client: TestClient) -> None:
    key = private_key()
    before = {f["id"] for f in client.get("/api/files").json()}
    response = client.post("/api/files", files={"file": ("svc.txt", f"historian service\n{key}".encode(), "text/plain")})
    assert response.status_code == 400 and "private_key" in response.json()["detail"]
    assert {f["id"] for f in client.get("/api/files").json()} == before
    assert key not in _audit_text() and key.splitlines()[1] not in _audit_text()


def test_upload_allow_stores_raises_and_notes_without_the_value(client: TestClient) -> None:
    key = aws_key()
    response = client.post(
        "/api/files",
        data={"classification": "normal"},
        files={"file": ("export.csv", f"tag,value\nkey,{key}\n".encode(), "text/csv")},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["classification"] == "restricted"
    assert any("content scanning found 1 aws_access_key" in note for note in body["quarantine_notes"])
    assert key not in json.dumps(body) and key not in _audit_text()


def test_a_document_is_scanned_from_its_text_not_its_bytes(client: TestClient) -> None:
    from docx import Document

    number = aadhaar()
    document = Document()
    document.add_paragraph(f"Permit holder Aadhaar {number}")
    path = get_config().settings.storage_root / "permit.docx"
    document.save(path)
    response = client.post("/api/files", data={"classification": "normal"}, files={
        "file": ("permit.docx", path.read_bytes(),
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 201, response.text
    assert response.json()["classification"] == "sensitive"
    assert number not in _audit_text() and number.replace(" ", "") not in _audit_text()


def test_knowledge_ingest_refuses_blocked_content_and_raises_marked_content(client: TestClient, monkeypatch) -> None:
    from backend.api.routes import system as system_routes

    calls: list[Sensitivity] = []

    class Index:
        """Records what it was asked to index; never embeds, so no model is called."""

        async def ingest_file(self, path, *, department, classification, version):
            calls.append(classification)
            raise ValueError("the fake index records the classification only")

    monkeypatch.setattr(system_routes, "get_knowledge_base", lambda: Index())
    key = private_key()
    response = client.post("/api/knowledge/documents", data={"classification": "normal"},
                           files={"file": ("keys.md", f"# Keys\n\n{key}".encode(), "text/markdown")})
    assert response.status_code == 400 and "private_key" in response.json()["detail"] and calls == []
    assert key.splitlines()[1] not in _audit_text()

    client.post("/api/knowledge/documents", data={"classification": "normal"},
                files={"file": ("memo.md", b"**Classification:** Restricted\n\nDesign basis.", "text/markdown")})
    assert calls == [Sensitivity.RESTRICTED]
